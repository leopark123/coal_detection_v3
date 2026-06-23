"""
将 Markdown 文件转换为 Word 文档（.docx）
用法：python tools/md_to_docx.py docs/部署手册.md
"""

import sys
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_table_border(table):
    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "AAAAAA")
        tblBorders.append(el)
    tblPr.append(tblBorders)


def add_code_block(doc, lines: list[str]):
    """灰底等宽字体代码块"""
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), "F5F5F5")
        p._p.get_or_add_pPr().append(shading)
        run = p.add_run(line if line else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def apply_inline(run_text: str, para):
    """处理行内 `code` 和 **bold**"""
    pattern = re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', run_text)
    for part in pattern:
        if part.startswith('`') and part.endswith('`'):
            run = para.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x40)
        elif part.startswith('**') and part.endswith('**'):
            run = para.add_run(part[2:-2])
            run.bold = True
        else:
            para.add_run(part)


def md_to_docx(md_path: str, out_path: str):
    md = Path(md_path).read_text(encoding="utf-8")
    lines = md.splitlines()

    doc = Document()

    # 页面设置
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    # 正文样式
    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    i = 0
    while i < len(lines):
        line = lines[i]

        # --- 代码块 ---
        if line.strip().startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            doc.add_paragraph()  # 空行
            add_code_block(doc, code_lines)
            doc.add_paragraph()  # 空行
            i += 1
            continue

        # --- 标题 ---
        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=1)
            p.runs[0].font.name = "微软雅黑"
            p.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_heading(line[3:].strip(), level=2)
            if p.runs:
                p.runs[0].font.name = "微软雅黑"
                p.runs[0].font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_heading(line[4:].strip(), level=3)
            if p.runs:
                p.runs[0].font.name = "微软雅黑"
            i += 1
            continue

        # --- 分割线 ---
        if re.match(r'^-{3,}$', line.strip()):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "BBBBBB")
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        # --- 表格 ---
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            # 过滤分隔行
            rows = [l for l in table_lines if not re.match(r'^\|[\s\-:|]+\|$', l)]
            if not rows:
                continue
            cols = [c.strip() for c in rows[0].strip("|").split("|")]
            table = doc.add_table(rows=len(rows), cols=len(cols))
            set_table_border(table)
            table.style = "Table Grid"
            for r_idx, row_line in enumerate(rows):
                cells = [c.strip() for c in row_line.strip("|").split("|")]
                for c_idx, cell_text in enumerate(cells):
                    if c_idx >= len(cols):
                        break
                    cell = table.rows[r_idx].cells[c_idx]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    apply_inline(cell_text, p)
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs:
                        run.font.name = "微软雅黑"
                        run.font.size = Pt(9.5)
                    if r_idx == 0:
                        set_cell_bg(cell, "E8F0FE")
                        for run in p.runs:
                            run.bold = True
            doc.add_paragraph()
            continue

        # --- 引用块 ---
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "12")
            left.set(qn("w:space"), "4")
            left.set(qn("w:color"), "4472C4")
            pBdr.append(left)
            pPr.append(pBdr)
            apply_inline(line[2:], p)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
                run.font.italic = True
            i += 1
            continue

        # --- 无序列表 ---
        if re.match(r'^[\-\*] ', line):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.5)
            apply_inline(line[2:], p)
            for run in p.runs:
                run.font.name = "微软雅黑"
                run.font.size = Pt(10)
            i += 1
            continue

        # --- 有序列表 ---
        if re.match(r'^\d+\. ', line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Cm(0.5)
            text = re.sub(r'^\d+\. ', '', line)
            apply_inline(text, p)
            for run in p.runs:
                run.font.name = "微软雅黑"
                run.font.size = Pt(10)
            i += 1
            continue

        # --- 空行 ---
        if line.strip() == "":
            i += 1
            continue

        # --- 普通段落 ---
        p = doc.add_paragraph()
        apply_inline(line, p)
        for run in p.runs:
            run.font.name = "微软雅黑"
            run.font.size = Pt(10.5)
        i += 1

    doc.save(out_path)
    print(f"已生成：{out_path}")


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "docs/部署手册.md"
    dst = sys.argv[2] if len(sys.argv) > 2 else str(Path(src).with_suffix(".docx"))
    md_to_docx(src, dst)
